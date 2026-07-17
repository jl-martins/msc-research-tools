#!/usr/bin/env -S uv run
import argparse
import logging
from pathlib import Path

import polars as pl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# File names
COMPARISON_FILE_NAME = "comparison.csv"

# Pattern Fills
POS_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
NEG_FILL = PatternFill(start_color="FCE5E5", end_color="FCE5E5", fill_type="solid")
NEUTRAL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")


# Logging
def setup_logging(verbose: bool = False) -> None:
    level = logging.DEBUG if verbose else logging.INFO
    logging.basicConfig(level=level, format="[%(levelname)s] %(message)s")
    logging.captureWarnings(True)


# Worksheet Logic
def populate_worksheet(ws: Worksheet, comparison_df: pl.DataFrame) -> None:
    ws.append(comparison_df.columns)
    for row in comparison_df.iter_rows():
        ws.append(row)


def apply_conditional_formatting(
    ws: Worksheet,
    headers: list[str],
    ci_low_col: str,
    ci_high_col: str,
) -> None:
    assert ci_low_col in headers
    assert ci_high_col in headers

    ci_low_letter = get_column_letter(headers.index(ci_low_col) + 1)
    ci_high_letter = get_column_letter(headers.index(ci_high_col) + 1)
    max_col_letter = get_column_letter(ws.max_column)

    target_range = f"A2:{max_col_letter}{ws.max_row}"
    rules = [
        FormulaRule(
            formula=[f"AND(${ci_low_letter}2>0, ${ci_high_letter}2>0)"],
            fill=POS_FILL,
            stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f"AND(${ci_low_letter}2<0, ${ci_high_letter}2<0)"],
            fill=NEG_FILL,
            stopIfTrue=True,
        ),
        FormulaRule(formula=["TRUE"], fill=NEUTRAL_FILL),
    ]

    for rule in rules:
        ws.conditional_formatting.add(target_range, rule)


def format_comparison(
    comparison_df: pl.DataFrame,
    ci_low_col: str,
    ci_high_col: str,
    csv_path: Path,
    output_path: Path,
) -> None:
    logger.info("Formatting '%s'...", csv_path)
    headers = comparison_df.columns
    missing_cols = {ci_low_col, ci_high_col} - set(headers)
    if missing_cols:
        logger.error("Missing columns: %s", ", ".join(sorted(missing_cols)))
        return

    wb = Workbook()
    ws = wb.active
    assert ws is not None

    ws.title = output_path.stem
    populate_worksheet(ws, comparison_df)
    apply_conditional_formatting(ws, headers, ci_low_col, ci_high_col)

    wb.save(output_path)
    logger.info("Formatted '%s' in '%s'", csv_path, output_path)


# Orchestration
def format_comparisons(root_dir: Path, effect_size: str) -> None:
    comparison_csvs = list(root_dir.rglob(COMPARISON_FILE_NAME))
    logger.info("Found %d comparison CSV(s)", len(comparison_csvs))

    ci_low_col = f"{effect_size}_CI_Low"
    ci_high_col = f"{effect_size}_CI_High"
    ci_schema = {
        ci_low_col: pl.Float64,
        ci_high_col: pl.Float64,
    }

    for csv_path in comparison_csvs:
        # noinspection PyBroadException
        try:
            comparison_df = pl.read_csv(csv_path, schema_overrides=ci_schema)
            output_path = csv_path.with_suffix(".xlsx")
            format_comparison(
                comparison_df,
                ci_low_col,
                ci_high_col,
                csv_path,
                output_path,
            )
        except Exception:
            logger.exception("Failed to format '%s'", csv_path)

    logger.info("Finished formatting comparison CSV(s)")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Format comparisons based on effect size confidence intervals",
        color=False,
    )
    parser.add_argument(
        "root_dir",
        nargs="?",
        default=Path.cwd(),
        type=Path,
        help="root directory for searching for comparisons (default: current working directory)",
    )
    parser.add_argument(
        "-s",
        "--effect-size",
        default="Cliffs_Delta",
        type=str,
        help="set the name of the effect size measure (default: Cliffs_Delta)",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="enable debug logging",
    )
    args = parser.parse_args()

    root_dir = args.root_dir.resolve()
    setup_logging(args.verbose)

    format_comparisons(root_dir, args.effect_size)


if __name__ == "__main__":
    main()
